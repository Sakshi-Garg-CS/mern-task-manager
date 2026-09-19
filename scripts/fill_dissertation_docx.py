"""Build Dissertation_Project-Report.docx from DISSERTATION_REPORT.txt with KU cover."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
import os

TXT = r"c:\Users\saksh\OneDrive\Desktop\TaskManager\DISSERTATION_REPORT.txt"
DOCX = r"c:\Users\saksh\OneDrive\Desktop\TaskManager\Dissertation_Project-Report.docx"
LOGO = r"c:\Users\saksh\OneDrive\Desktop\TaskManager\assets\ku-logo.png"

BLANK_RE = re.compile(r"^[_\s\.\-]+$")


def is_banner_line(line):
    s = line.strip()
    return len(s) > 10 and set(s) <= {"=", " "}


def parse_meta(lines):
    """Read key : value lines from the cover block in the txt file."""
    meta = {
        "report_title": "Web-Based Task Management System",
        "student_name": "<Student Name>",
        "roll_no": "<Roll No.>",
        "programme": "<Programme Name>",
        "course_code": "<Course Code>",
        "supervisor": "<Name of the Supervisor>",
        "chairperson": "<Name of the Chairperson>",
        "session": "<Session>",
        "submission_date": "<Date>",
        "supervisor_email": "",
        "chairperson_email": "chairperson.dcsa@kuk.ac.in",
        "chairperson_ext": "2133",
        "ca_reference": "CA/26/________",
        "supervisor_designation": "Professor",
        "organization": "Department of Computer Science and Applications, Kurukshetra University",
        "project_start": "<Start Date>",
        "project_end": "<End Date>",
        "degree_name": "",
    }
    key_map = {
        "report title": "report_title",
        "student name": "student_name",
        "enrollment no": "roll_no",
        "enrollment no.": "roll_no",
        "programme": "programme",
        "course code": "course_code",
        "guide / supervisor": "supervisor",
        "guide/supervisor": "supervisor",
        "session / year": "session",
        "session/year": "session",
        "date of submission": "submission_date",
        "chairperson": "chairperson",
        "chairperson / hod": "chairperson",
        "supervisor email": "supervisor_email",
        "chairperson email": "chairperson_email",
        "ca reference no": "ca_reference",
        "ca reference": "ca_reference",
        "supervisor designation": "supervisor_designation",
        "organization": "organization",
        "project start date": "project_start",
        "project end date": "project_end",
        "degree name": "degree_name",
    }
    for line in lines:
        if ":" not in line:
            continue
        left, right = line.split(":", 1)
        key = left.strip().lower()
        val = right.strip()
        if key not in key_map:
            continue
        if val and not BLANK_RE.match(val):
            meta[key_map[key]] = val
    if not meta["degree_name"] or meta["degree_name"].startswith("<"):
        meta["degree_name"] = meta["programme"]
    return meta


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.15):
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(int(before * 20)))
    spacing.set(qn("w:after"), str(int(after * 20)))
    spacing.set(qn("w:line"), str(int(line * 240)))
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)


def add_centered(
    doc, text, *, bold=False, italic=False, size=12, space_before=0, space_after=6
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    set_paragraph_spacing(p, before=space_before, after=space_after)
    return p


def remove_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)


def add_cover_page(doc, meta):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    add_centered(doc, "Project Report", bold=True, size=16, space_after=4)
    add_centered(doc, "On", size=14, space_after=4)
    add_centered(doc, meta["report_title"], bold=True, size=16, space_after=10)

    add_centered(
        doc,
        "Submitted in the partial fulfillment of the requirement",
        size=12,
        space_after=2,
    )
    add_centered(doc, "for the award of degree of", size=12, space_after=6)
    add_centered(doc, meta["programme"], bold=True, size=14, space_after=4)
    add_centered(doc, f"(Session: {meta['session']})", bold=True, size=12, space_after=14)

    if os.path.isfile(LOGO):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_logo.add_run()
        run.add_picture(LOGO, width=Inches(2.2))
        set_paragraph_spacing(p_logo, before=6, after=18)
    else:
        add_centered(doc, "[Insert Kurukshetra University logo]", italic=True, size=11, space_after=18)

    table = doc.add_table(rows=1, cols=2)
    remove_table_borders(table)
    table.autofit = False
    table.columns[0].width = Inches(3.25)
    table.columns[1].width = Inches(3.25)

    left, right = table.rows[0].cells

    def fill_cell(cell, lines):
        cell.text = ""
        for i, (text, bold, size) in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run.bold = bold
            set_paragraph_spacing(p, after=4)

    fill_cell(
        left,
        [
            ("Supervised By:", True, 12),
            (meta["supervisor"], False, 12),
        ],
    )
    fill_cell(
        right,
        [
            ("Submitted By:", True, 12),
            (meta["student_name"], False, 12),
            (f"Roll No. {meta['roll_no']}", False, 12),
        ],
    )

    doc.add_paragraph()
    add_centered(
        doc,
        "Department of Computer Science and Applications",
        bold=True,
        size=14,
        space_before=24,
        space_after=6,
    )
    add_centered(
        doc,
        "Kurukshetra University, Kurukshetra – 136119",
        bold=True,
        size=14,
        space_after=0,
    )

    doc.add_page_break()


def add_left_body(doc, text, *, size=12, space_after=10, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    set_paragraph_spacing(p, before=space_before, after=space_after, line=1.5)
    return p


def add_declaration_page(doc, meta):
    add_centered(
        doc,
        "Declaration by the Candidate",
        bold=True,
        size=14,
        space_before=6,
        space_after=18,
    )

    title = meta["report_title"]
    p1 = (
        f"I, {meta['student_name']}, bearing Roll Number {meta['roll_no']}, "
        f"hereby declare that the project report titled \u201c{title}\u201d, "
        f"submitted in partial fulfillment of the requirements for the award of "
        f"the {meta['programme']} of the course {meta['course_code']}, is my "
        f"original work and has been carried out under the guidance of "
        f"{meta['supervisor']}."
    )
    p2 = (
        "I also declare that this work has not been submitted previously by me "
        "or any other individual for the award of any degree, diploma, or "
        "certificate at this or any other institution."
    )
    p3 = (
        "Any references and resources used during the development of this project "
        "have been duly acknowledged in the report."
    )

    add_left_body(doc, p1, space_after=12)
    add_left_body(doc, p2, space_after=12)
    add_left_body(doc, p3, space_after=24)
    add_left_body(doc, "(Candidate\u2019s Signature)", space_before=36, space_after=0)

    doc.add_page_break()


def _font_run(run, *, bold=False, size=12, underline=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    if underline:
        run.underline = WD_UNDERLINE.SINGLE


def _para_in_cell(cell, align=WD_ALIGN_PARAGRAPH.CENTER):
    return cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()


def add_dcsa_letterhead(doc):
    """KU DCSA header: logo top-left, department/university lines centered."""
    table = doc.add_table(rows=1, cols=2)
    remove_table_borders(table)
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.15)
    logo_cell, text_cell = table.rows[0].cells

    if os.path.isfile(LOGO):
        p_logo = _para_in_cell(logo_cell, WD_ALIGN_PARAGRAPH.LEFT)
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p_logo.add_run()
        run.add_picture(LOGO, width=Inches(1.15))

    lines = [
        ("Department of Computer Science and Applications", True, 12),
        ("KURUKSHETRA UNIVERSITY KURUKSHETRA", True, 14),
        ("(Established by the State Legislature Act XII of 1956)", False, 11),
        ("('A++' Grade, NAAC Accredited)", False, 11),
    ]
    for i, (text, bold, size) in enumerate(lines):
        p = _para_in_cell(text_cell) if i == 0 else text_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        _font_run(run, bold=bold, size=size)
        set_paragraph_spacing(p, after=2)


def add_horizontal_rule(doc, *, thick=False):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "24" if thick else "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    set_paragraph_spacing(p, after=12)


def format_supervisor_name(name):
    if name.strip().startswith("<"):
        return name
    n = name.strip()
    if not n.lower().startswith("dr"):
        n = f"Dr. {n}"
    return n


def format_chairperson_name(name):
    if name.strip().startswith("<"):
        return name
    n = name.strip()
    if not n.lower().startswith("prof"):
        n = f"Prof. {n}"
    return n


def _cell_para(cell, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    return p


def add_justified_body(doc, text, *, space_after=12, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    _font_run(run, size=12)
    set_paragraph_spacing(p, before=space_before, after=space_after, line=1.5)
    return p


def add_supervisor_certificate_page(doc, meta):
    add_dcsa_letterhead(doc)

    bar = doc.add_table(rows=1, cols=2)
    remove_table_borders(bar)
    bar.autofit = False
    bar.columns[0].width = Inches(3.5)
    bar.columns[1].width = Inches(3.0)
    left_cell, right_cell = bar.rows[0].cells

    sup_name = format_supervisor_name(meta["supervisor"])
    p_left = _para_in_cell(left_cell, WD_ALIGN_PARAGRAPH.LEFT)
    p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_left = p_left.add_run(sup_name)
    _font_run(run_left, size=12, underline=True)

    email = meta.get("supervisor_email") or "___________________________"
    dated = meta.get("submission_date") or "dd/mm/2026"
    if dated.startswith("<"):
        dated = "dd/mm/2026"

    p_right = _para_in_cell(right_cell, WD_ALIGN_PARAGRAPH.LEFT)
    p_right.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p_right.add_run(f"Email: {email}\n")
    _font_run(r1, size=12)
    r2 = p_right.add_run(f"Dated: {dated}")
    _font_run(r2, size=12)
    set_paragraph_spacing(p_left, after=6)
    set_paragraph_spacing(p_right, after=6)

    add_horizontal_rule(doc, thick=True)

    add_centered(
        doc,
        "Certificate by the Supervisor",
        bold=True,
        size=14,
        space_before=6,
        space_after=16,
    )

    title_q = meta["report_title"]
    session = meta["session"]
    p1 = (
        f"This is to certify that the project report titled \u201c{title_q}\u201d "
        f"submitted by {meta['student_name']}, Enrollment Number {meta['roll_no']}, "
        f"in partial fulfillment of the requirements for the award of the "
        f"{meta['programme']} of the course {meta['course_code']}, is a bonafide "
        f"record of original work carried out by him/her under my supervision during "
        f"the academic session {session}."
    )
    p2 = (
        "This work has not been submitted previously for the award of any degree, "
        "diploma, or certificate in this or any other university/institute."
    )

    add_justified_body(doc, p1, space_after=14)
    add_justified_body(doc, p2, space_after=36)
    add_left_body(doc, "(Supervisor\u2019s Signature)", space_before=12, space_after=0)

    doc.add_page_break()


def add_chairperson_certificate_page(doc, meta):
    """DCSA chairperson certificate — official letterhead layout."""
    header = doc.add_table(rows=1, cols=3)
    remove_table_borders(header)
    header.autofit = False
    header.columns[0].width = Inches(1.25)
    header.columns[1].width = Inches(3.75)
    header.columns[2].width = Inches(1.5)
    logo_cell, mid_cell, right_cell = header.rows[0].cells

    if os.path.isfile(LOGO):
        p_logo = _cell_para(logo_cell, WD_ALIGN_PARAGRAPH.LEFT)
        p_logo.add_run().add_picture(LOGO, width=Inches(1.1))

    for i, (text, bold, size) in enumerate(
        [
            ("Department of Computer Science and Applications", True, 12),
            ("KURUKSHETRA UNIVERSITY KURUKSHETRA", True, 14),
            ("(Established by the State Legislature Act XII of 1956)", False, 11),
            ("('A++' Grade, NAAC Accredited)", False, 11),
        ]
    ):
        p = _cell_para(mid_cell, WD_ALIGN_PARAGRAPH.CENTER) if i == 0 else mid_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        _font_run(run, bold=bold, size=size)
        set_paragraph_spacing(p, after=2)

    ext = meta.get("chairperson_ext", "2133")
    email = meta.get("chairperson_email", "chairperson.dcsa@kuk.ac.in")
    ca_no = meta.get("ca_reference", "CA/26/________")
    dated = meta.get("submission_date") or "dd/mm/2026"
    if str(dated).startswith("<"):
        dated = "dd/mm/2026"

    p_right = _cell_para(right_cell, WD_ALIGN_PARAGRAPH.LEFT)
    for line in (f"Ext. {ext}", f"Email: {email}", f"No.{ca_no}", f"Dated: {dated}"):
        run = p_right.add_run(line + "\n")
        _font_run(run, size=11)
    set_paragraph_spacing(p_right, after=4)

    bar = doc.add_table(rows=1, cols=2)
    remove_table_borders(bar)
    bar.columns[0].width = Inches(4.0)
    bar.columns[1].width = Inches(2.5)
    left_cell, _ = bar.rows[0].cells

    chair_name = format_chairperson_name(meta["chairperson"])
    p_ch = _cell_para(left_cell)
    r1 = p_ch.add_run(chair_name + "\n")
    _font_run(r1, bold=True, size=12)
    r2 = p_ch.add_run("Chairperson")
    _font_run(r2, bold=True, size=12)
    set_paragraph_spacing(p_ch, after=6)

    add_horizontal_rule(doc, thick=True)

    title_q = meta["report_title"]
    supervisor = format_supervisor_name(meta["supervisor"])
    designation = meta.get("supervisor_designation", "Professor")

    p1 = (
        f"This is to certify that the project report entitled \u201c{title_q}\u201d "
        f"is a bona fide work carried out by {meta['student_name']}, Roll No. "
        f"{meta['roll_no']}, student of {meta['programme']}, under the guidance and "
        f"supervision of {supervisor}, {designation}, Department of Computer Science "
        f"and Applications, Kurukshetra University, Kurukshetra."
    )
    org = meta.get("organization", "Department of Computer Science and Applications, Kurukshetra University")
    p2 = (
        f"The project work was conducted at {org} during the period from "
        f"{meta.get('project_start', '<Start Date>')} to {meta.get('project_end', '<End Date>')}. "
        f"This project is approved for submission toward the partial fulfillment of the "
        f"requirements for the award of the degree of {meta['degree_name']}."
    )

    add_justified_body(doc, p1, space_after=14, space_before=8)
    add_justified_body(doc, p2, space_after=28)

    sign_date = meta.get("submission_date") or "<Current Date>"
    for line in (
        "(Signature)",
        chair_name,
        "Chairperson, Department of Computer Science and Applications,",
        "Kurukshetra University, Kurukshetra.",
        f"Date: {sign_date}",
        "Seal:",
    ):
        add_left_body(doc, line, space_after=4)

    doc.add_page_break()


def add_paragraph(doc, text, style=None, bold=False, italic=False, size=12):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def section_start_index(lines, section_name):
    """Find banner line before a section heading (e.g. ABSTRACT)."""
    target = section_name.upper()
    for i, line in enumerate(lines):
        if target not in line.strip().upper():
            continue
        start = i
        while start > 0 and (
            not lines[start - 1].strip() or is_banner_line(lines[start - 1])
        ):
            start -= 1
        if start > 0 and is_banner_line(lines[start - 1]):
            return start - 1
        return start
    return 0


def main():
    with open(TXT, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    meta = parse_meta(lines)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    add_cover_page(doc, meta)
    add_declaration_page(doc, meta)
    add_supervisor_certificate_page(doc, meta)
    add_chairperson_certificate_page(doc, meta)

    start = section_start_index(lines, "ABSTRACT")
    in_screenshot_block = False
    skip_paren_note = False

    for line in lines[start:]:
        raw = line.rstrip()
        stripped = raw.strip()

        if not stripped:
            if in_screenshot_block:
                in_screenshot_block = False
            continue

        if skip_paren_note:
            if stripped.startswith("("):
                continue
            skip_paren_note = False

        if stripped.startswith("(University/Department"):
            continue

        if is_banner_line(stripped):
            title = stripped.replace("=", "").strip()
            if title.upper() == "COVER PAGE (AUTO IN WORD)":
                skip_paren_note = True
                continue
            if "DECLARATION" in title.upper() and "AUTO" in title.upper():
                skip_paren_note = True
                continue
            if "CERTIFICATE" in title.upper() and "AUTO" in title.upper():
                skip_paren_note = True
                continue
            add_paragraph(doc, title or " ", bold=True, size=14)
            continue

        if stripped.startswith("CHAPTER ") or stripped == "END OF REPORT":
            add_paragraph(doc, stripped, bold=True, size=14)
            continue

        if re.match(r"^\d+\.\d+\s", stripped):
            add_paragraph(doc, stripped, bold=True, size=12)
            continue

        if stripped.startswith("APPENDIX ") or stripped.startswith("Figure "):
            add_paragraph(doc, stripped, bold=True, size=11)
            continue

        if "[SCREENSHOT PLACEHOLDER" in stripped or stripped.startswith("|  FIGURE"):
            in_screenshot_block = True
            add_paragraph(doc, stripped.lstrip("| "), italic=True, size=11)
            continue

        if in_screenshot_block and (stripped.startswith("|") or stripped.startswith("[")):
            add_paragraph(doc, stripped.lstrip("| "), italic=True, size=11)
            continue

        if stripped.startswith("------"):
            in_screenshot_block = False
            continue

        if stripped.startswith("+---") or stripped.startswith("| ID"):
            add_paragraph(doc, stripped, size=10)
            continue

        # Skip duplicate meta lines inside body
        if re.match(
            r"^(Report Title|Student Name|Enrollment|Programme|Course Code|Guide|Chairperson|Supervisor Email|Supervisor Designation|Organization|Project Start|Project End|Degree Name|CA Reference|Session|Date of)",
            stripped,
        ):
            continue

        add_paragraph(doc, stripped, size=12)

    doc.save(DOCX)
    print(f"Saved: {DOCX}")


if __name__ == "__main__":
    main()
